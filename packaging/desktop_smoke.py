"""Run the packaged desktop executable through its self-test and check it opens no sockets.

Windows only. Usage:
    uv run --group build python packaging/desktop_smoke.py dist/MarkItDown-Web.exe

The executable converts one sample of each main format in its real window (see
app/selftest.py). Meanwhile this script polls netstat for sockets owned by the app's
process tree, WebView2 helpers included: there must be none.
"""

from __future__ import annotations

import csv
import io
import json
import os
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

MARKER = "SmokeMarker"


def _pdf() -> bytes:
    stream = f"BT /F1 24 Tf 72 720 Td ({MARKER} pdf) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (number, body))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1, xref))
    return out.getvalue()


def _docx() -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" '
            'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            "word/document.xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body><w:p><w:r><w:t>{MARKER} docx</w:t></w:r></w:p></w:body></w:document>",
        )
    return out.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active["A1"] = f"{MARKER} xlsx"
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = f"{MARKER} pptx"
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


SAMPLES = {
    "sample.pdf": _pdf,
    "sample.docx": _docx,
    "sample.xlsx": _xlsx,
    "sample.pptx": _pptx,
    "sample.html": lambda: f"<html><body><h1>{MARKER} html</h1></body></html>".encode(),
    "sample.csv": lambda: f"name,value\n{MARKER} csv,1\n".encode(),
}


def _process_tree(root: int) -> dict[int, str]:
    listing = subprocess.run(
        ["powershell", "-NoProfile", "-Command",
         "Get-CimInstance Win32_Process | Select-Object ProcessId,ParentProcessId,Name"
         " | ConvertTo-Csv -NoTypeInformation"],
        capture_output=True, text=True, check=True,
    ).stdout
    rows = list(csv.DictReader(io.StringIO(listing)))
    children: dict[int, list[tuple[int, str]]] = {}
    for row in rows:
        children.setdefault(int(row["ParentProcessId"]), []).append(
            (int(row["ProcessId"]), row["Name"]))
    tree, pending = {}, [root]
    while pending:
        pid = pending.pop()
        for child, name in children.get(pid, []):
            if child not in tree:
                tree[child] = name
                pending.append(child)
    return tree


def _sockets(pids: dict[int, str]) -> set[tuple[str, str]]:
    output = subprocess.run(["netstat", "-ano"], capture_output=True, text=True).stdout
    found = set()
    for line in output.splitlines():
        parts = line.split()
        if parts and parts[0] in {"TCP", "UDP"} and parts[-1].isdigit():
            pid = int(parts[-1])
            if pid in pids:
                found.add((pids[pid], " ".join(parts[:-1])))
    return found


def main(executable: str) -> int:
    folder = Path(tempfile.mkdtemp(prefix="markitdown-smoke-"))
    for name, build in SAMPLES.items():
        (folder / name).write_bytes(build())
    env = dict(os.environ, MARKITDOWN_DESKTOP_SELFTEST=str(folder))
    process = subprocess.Popen([executable], env=env)
    sockets: set[tuple[str, str]] = set()
    names: set[str] = set()
    deadline = time.monotonic() + 300
    while process.poll() is None and time.monotonic() < deadline:
        tree = _process_tree(process.pid)
        tree[process.pid] = Path(executable).name
        names.update(tree.values())
        sockets |= _sockets(tree)
        time.sleep(0.5)
    if process.poll() is None:
        process.kill()
        print("FAIL the executable did not finish its self-test in time", file=sys.stderr)
        return 1

    print("processes:", ", ".join(sorted(names)))
    failures = []
    for owner, socket in sorted(sockets):
        failures.append(f"{owner} opened a socket: {socket}")

    result_path = folder / "result.json"
    if not result_path.exists():
        failures.append(f"no result.json (exit code {process.returncode})")
    else:
        result = json.loads(result_path.read_text(encoding="utf-8"))
        if result.get("error"):
            failures.append(f"self-test error: {result['error']}")
        if not result.get("ready"):
            failures.append("the page or the Python bridge never became ready")
        for name in SAMPLES:
            state = result.get("samples", {}).get(name)
            if not state:
                failures.append(f"{name}: not run")
            elif MARKER not in state["markdown"] or MARKER not in state["preview"]:
                failures.append(f"{name}: {state['state']!r} {state['meta']!r}")
            else:
                print(f"{name} ok ({state['meta']})")

    for failure in failures:
        print(f"FAIL {failure}", file=sys.stderr)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
