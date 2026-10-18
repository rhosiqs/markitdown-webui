"""Start a built executable and convert one file of each main format through its HTTP API.

Usage: uv run --group build python packaging/smoke_test.py dist/MarkItDown-Web.exe
"""

from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import time
import urllib.error
import urllib.request
import uuid
import zipfile

PORT = 18765
BASE = f"http://127.0.0.1:{PORT}"
MARKER = "SmokeMarker"


def _pdf() -> bytes:
    stream = f"BT /F1 24 Tf 72 720 Td ({MARKER} pdf) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (number, body))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1, xref))
    return out.getvalue()


def _docx() -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" '
            'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            "word/document.xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body><w:p><w:r><w:t>{MARKER} docx</w:t></w:r></w:p></w:body></w:document>",
        )
    return out.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active["A1"] = f"{MARKER} xlsx"
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = f"{MARKER} pptx"
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


SAMPLES = {
    "sample.pdf": _pdf,
    "sample.docx": _docx,
    "sample.xlsx": _xlsx,
    "sample.pptx": _pptx,
    "sample.html": lambda: f"<html><body><h1>{MARKER} html</h1></body></html>".encode(),
    "sample.csv": lambda: f"name,value\n{MARKER} csv,1\n".encode(),
}


def _request(path: str, body: bytes | None = None, content_type: str | None = None) -> bytes:
    request = urllib.request.Request(BASE + path, data=body)
    if content_type:
        request.add_header("Content-Type", content_type)
    with urllib.request.urlopen(request, timeout=120) as response:
        return response.read()


def _convert(name: str, data: bytes) -> str:
    boundary = uuid.uuid4().hex
    body = (
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; "
        f"filename=\"{name}\"\r\nContent-Type: application/octet-stream\r\n\r\n"
    ).encode() + data + f"\r\n--{boundary}--\r\n".encode()
    payload = _request("/api/convert", body, f"multipart/form-data; boundary={boundary}")
    return json.loads(payload)["markdown"]


def _wait_for_health(process: subprocess.Popen, timeout: float = 180.0) -> None:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if process.poll() is not None:
            raise RuntimeError(f"The executable exited early with code {process.returncode}.")
        try:
            if json.loads(_request("/api/health")) == {"status": "ok"}:
                return
        except (urllib.error.URLError, ConnectionError, TimeoutError):
            time.sleep(1)
    raise RuntimeError("The executable did not answer /api/health in time.")


def main(executable: str) -> int:
    env = dict(os.environ, MARKITDOWN_WEB_PORT=str(PORT), MARKITDOWN_WEB_OPEN_BROWSER="0")
    process = subprocess.Popen([executable], env=env, stdin=subprocess.DEVNULL)
    failures = []
    try:
        started = time.monotonic()
        _wait_for_health(process)
        print(f"health ok after {time.monotonic() - started:.1f}s")

        page = _request("/").decode()
        assert "/static/app.js" in page, "index page does not reference app.js"
        _request("/static/app.js")
        _request("/static/styles.css")
        print("index and static files ok")

        rendered = json.loads(_request(
            "/api/render", json.dumps({"markdown": "# Title"}).encode(), "application/json"))
        assert "<h1>Title</h1>" in rendered["html"], rendered
        print("render ok")

        for name, build in SAMPLES.items():
            try:
                markdown = _convert(name, build())
            except urllib.error.HTTPError as exc:
                failures.append(f"{name}: HTTP {exc.code} {exc.read()[:300]!r}")
                continue
            if MARKER not in markdown:
                failures.append(f"{name}: marker missing from {markdown[:200]!r}")
            else:
                print(f"{name} ok")
    finally:
        process.terminate()
        try:
            process.wait(timeout=15)
        except subprocess.TimeoutExpired:
            process.kill()

    for failure in failures:
        print(f"FAIL {failure}", file=sys.stderr)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
